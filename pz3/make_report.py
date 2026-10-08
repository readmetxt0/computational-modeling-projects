import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from training_matrix import split_textures, build_Y, brightness, M, N, n
from binary_matrix import tolerances, binarize, DELTA
from reference_vectors import reference_vectors, RO

imgs = split_textures("images/textures_var10.png")
Y = build_Y(imgs)
ndk, vdk = tolerances(Y, DELTA)
X = binarize(Y, ndk, vdk)
EV = reference_vectors(X, RO)

# ---------- рисунки ----------
fig, ax = plt.subplots(1, 2, figsize=(7, 2))
for c in range(M):
    ax[c].imshow(np.tile(EV[c][None, :], (10, 1)), cmap="gray", vmin=0, vmax=1, aspect="auto")
    ax[c].set_title(f"Эталонный вектор класса {c+1} ({int(EV[c].sum())} из {N} единиц)", fontsize=8)
    ax[c].set_yticks([]); ax[c].set_xlabel("признак i")
plt.tight_layout(); plt.savefig("images/ev.png", dpi=150); plt.close()

# ---------- помощники для docx ----------
def new_doc():
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"; d.styles["Normal"].font.size = Pt(14)
    for s in d.sections:
        s.left_margin = Cm(3); s.right_margin = Cm(1.5); s.top_margin = s.bottom_margin = Cm(2)
    return d

def P(d, text, bold=False, center=False, italic=False):
    par = d.add_paragraph(); r = par.add_run(text); r.bold = bold; r.italic = italic
    if center: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return par

def CODE(d, text):
    par = d.add_paragraph(); r = par.add_run(text); r.font.name = "Courier New"; r.font.size = Pt(9)

def PIC(d, path, cap, w=11):
    d.add_picture(path, width=Cm(w)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(d, cap, center=True)

def QA(d, items):
    P(d, "Ответы на контрольные вопросы", True)
    for i, (q, a) in enumerate(items, 1):
        P(d, f"{i}. {q}", True); P(d, a)

def head(d, num, title, goal):
    P(d, f"Практическое занятие {num}", True, True); P(d, title, True, True); P(d, "Вариант 10", italic=True, center=True)
    P(d, "Цель работы: ", True).add_run(goal).bold = False

def frag(arr, r=8, c=20):
    return "\n".join(" ".join(f"{int(v):3d}" if arr.max() > 1 else str(int(v)) for v in row[:c]) for row in arr[:r])


# ================= ПЗ-3 =================
d = new_doc()
head(d, 3, "Определение эталонных геометрических векторов классов распознавания",
     "научиться формировать геометрические центры классов распознавания.")
P(d, "1. Краткие теоретические сведения", True)
P(d, "Эталонный вектор x_m – математическое ожидание реализаций класса X_m^0 – детерминированный структурированный "
     "бинарный вектор x_m = <x_m,1, …, x_m,i, …, x_m,N>, m = 1..M, где x_m,i = 1, если значение i-го признака находится "
     "в нормированном поле допусков, и 0 – если не находится. При обосновании гипотезы компактности реализаций образа "
     "геометрическим центром класса является бинарный эталонный вектор x_m.")
P(d, "Элементы эталонного вектора: EV[k,i] = 1, если (1/n)·Σ_j X[k,i,j] > 0,5, и EV[k,i] = 0, если (1/n)·Σ_j X[k,i,j] ≤ 0,5.")
P(d, "2. Ход работы", True)
P(d, "Для каждого класса k и признака i вычисляется среднее значение бинарных признаков по n = 100 реализациям "
     "и сравнивается с уровнем селекции ρ = 0,5. Получен массив EV[1..2, 1..100].")
P(d, "3. Результаты", True)
PIC(d, "images/ev.png", "Рисунок 1 – Эталонные векторы классов распознавания", 14)
P(d, f"Число единиц: EV[1] – {int(EV[0].sum())}, EV[2] – {int(EV[1].sum())} из {N}. Кодовое расстояние между центрами классов "
     f"= {int((EV[0] != EV[1]).sum())}.")
for c in range(M):
    P(d, f"Эталонный вектор класса {c+1}:", italic=True)
    CODE(d, "".join(str(int(v)) for v in EV[c]))
P(d, "4. Программная реализация", True)
CODE(d, open("reference_vectors.py", encoding="utf-8").read())
QA(d, [
 ("Что называется эталонным вектором класса распознавания?",
  "Математическое ожидание реализаций класса – бинарный вектор x_m = <x_m,1, …, x_m,N>, представляющий геометрический центр класса."),
 ("Как формируются элементы двоичного эталонного вектора?",
  "x_m,i = 1, если среднее по реализациям значение бинарного признака превышает уровень селекции ρ_m (по умолчанию 0,5), иначе 0."),
 ("Что является геометрическим центром класса распознавания?",
  "Бинарный эталонный вектор x_m (при гипотезе компактности реализаций образа)."),
 ("Какой компонент использован для корректного отображения геометрических центров классов распознавания?",
  "В методичке – PictureBox/Bitmap (C#) для рисования вектора полосой чёрных и белых пикселей; в работе – imshow (matplotlib) "
  "и текстовое представление вектора."),
])
P(d, "Вывод", True)
P(d, "Сформирован массив эталонных векторов EV[1..2, 1..100]. Класс 2 имеет эталонный вектор из единиц (его реализации плотнее "
     "попадают в поля допусков), у класса 1 единиц меньше; центры различаются на "
     f"{int((EV[0] != EV[1]).sum())} разрядов, что позволяет строить матрицу кодовых расстояний (ПЗ-4).")
d.save("PZ3_otchet.docx")
